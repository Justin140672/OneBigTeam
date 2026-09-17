from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

OUT = r"C:\Users\justi\source\repos\OneBigTeam\One Big Team Technical Architecture.docx"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def border(cell):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); b.append(e)
    tcPr.append(b)
def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'1F4E78'); border(c)
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
    trPr = t.rows[0]._tr.get_or_add_trPr(); tblHeader = OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'), 'true'); trPr.append(tblHeader)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); border(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2==1: shade(cells[i],'EEF4F8')
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
                for r in p.runs: r.font.size=Pt(8.5)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(4)
    return t
def add_heading(doc, text, level=1):
    p=doc.add_heading(text, level); p.paragraph_format.space_before=Pt(14 if level==1 else 10); p.paragraph_format.space_after=Pt(6); return p
def add_p(doc, text, boldlead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.12
    if boldlead and text.startswith(boldlead):
        r=p.add_run(boldlead); r.bold=True; p.add_run(text[len(boldlead):])
    else: p.add_run(text)
    return p
def bullets(doc, items):
    for x in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.add_run(x)

d=Document()
sec=d.sections[0]; sec.top_margin=Inches(.8); sec.bottom_margin=Inches(.8); sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
styles=d.styles
styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos'); styles['Normal'].font.size=Pt(10)
for name,size in [('Title',24),('Heading 1',16),('Heading 2',12),('Heading 3',10.5)]:
    s=styles[name]; s.font.name='Aptos Display' if name!='Normal' else 'Aptos'; s._element.rPr.rFonts.set(qn('w:eastAsia'),s.font.name); s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)

p=d.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('One Big Team Technical Architecture')
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Architecture, operations and deployment guide'); r.italic=True; r.font.size=Pt(12)
d.add_paragraph('Version 1.0 | 17 September 2026', style='Subtitle').alignment=WD_ALIGN_PARAGRAPH.CENTER
d.add_paragraph()
add_p(d,'This document describes the technical architecture of One Big Team, a multi-tenant HR platform for small and medium-sized organisations. It explains the application structure, hosting model, security boundaries and operational controls. The appendices provide the practical setup steps for Cloudflare, Postmark, Railway and the GitHub Actions jobs that operate the delivery pipeline.')

add_heading(d,'Document Review')
add_heading(d,'Document Approvers',2)
table(d,['Name','Role','Mandatory'],[['Product owner','Product and operational approval','Yes'],['Technical lead','Architecture and security approval','Yes'],['Operations owner','Hosting and recovery approval','Yes']],[1.55,3.55,1.2])
add_heading(d,'Document Sign Off',2)
table(d,['Name','Role','Review date'],[['','',''],['','',''],['','','']],[1.8,3.3,1.2])
add_heading(d,'Version History',2)
table(d,['Version','Author','Date','Changes'],[['1.0','One Big Team','17 September 2026','Initial architecture and operational setup guide']],[.8,1.6,1.4,3.2])

add_heading(d,'Introduction')
add_p(d,'The purpose of this document is to describe the architecture of the One Big Team solution. It is intended for product, engineering and operations stakeholders who need to understand how the platform is built, hosted, secured and changed safely.')
add_p(d,'It describes:')
bullets(d,['A general description of the platform and its users.','The logical architecture, layers and principal components.','The physical architecture and managed services on which the platform runs.','The technical choices and the operational controls that support them.','The deployment pipeline, recovery approach and background-job model.'])
add_heading(d,'Abbreviations and Glossary',2)
table(d,['Term','Meaning'],[['API','Application Programming Interface'],['CI','Continuous Integration'],['CD','Continuous Delivery'],['DNS','Domain Name System'],['SLO','Service Level Objective'],['WAF','Web Application Firewall'],['Hangfire','The .NET background-processing framework used by the API']],[1.35,5.75])

add_heading(d,'Architecture')
add_heading(d,'Architecture overview',2)
add_p(d,'One Big Team is a modular monolith for HR administration and employee self-service. It supports organisations in managing people, leave, sickness, recruitment, onboarding, offboarding, probation, tasks, documents, assets, notifications, reporting and support. The primary users are company administrators, HR teams, managers and employees; a separate administrative experience supports platform operations.')
add_p(d,'The platform prioritises a straightforward operational model: a browser-based user experience, a FastEndpoints API, a PostgreSQL-backed modular domain model and managed specialist services for authentication, storage and email. It avoids microservices and Kubernetes in version 1 so that the system remains economical and supportable for SME customers.')
add_heading(d,'Physical architecture overview',2)
add_p(d,'Production traffic enters through Cloudflare, which provides authoritative DNS, edge TLS and optional web protection. Cloudflare routes public HTTP traffic to Railway. Railway hosts the API and web-facing services, builds from the selected Git commit and exposes each service through a public domain. The API connects to Supabase for PostgreSQL, authentication and private object storage, and calls Postmark for transactional email delivery.')
table(d,['Layer','Technology','Responsibility'],[['Edge','Cloudflare','DNS, TLS, traffic proxying, redirect and web protection'],['Hosting','Railway','Deploys and runs API, application, marketing and admin services'],['Application','ASP.NET Core, Blazor, FastEndpoints','Browser UI, APIs, authentication hand-off and business workflows'],['Data','Supabase PostgreSQL','Tenant data, module schemas, audit data and Hangfire storage'],['Identity and files','Supabase Auth and Storage','User identity, tokens, private documents and signed file access'],['Email','Postmark','Transactional notices, invitations, reminders and contact-form delivery'],['Local development','.NET Aspire','Local service orchestration, diagnostics and observability']],[1.25,1.75,4.1])
add_heading(d,'Logical architecture overview',2)
add_p(d,'The solution follows a modular-monolith design. Each business area owns its domain behaviour and persistence boundary while sharing a consistent host, infrastructure and security model. The browser applications communicate with the API over HTTPS; business modules do not send email directly, write to shared infrastructure directly or bypass tenant and authorisation rules.')
table(d,['Component','Purpose'],[['HR.Web','The employee and manager web experience.'],['HR.Marketing','Public marketing site, including the controlled contact-form relay.'],['HR.Admin.Web','Administrative operations interface.'],['HR.Api','FastEndpoints host, API composition, migration orchestration, health endpoints and Hangfire server.'],['Business modules','Vertical business capabilities such as employees, leave, documents, notifications and recruitment.'],['Infrastructure','Cross-cutting persistence, email, background jobs, security, auditing and storage adapters.'],['Shared kernel','Common domain concepts, validation and idempotency primitives.']],[1.7,5.4])
add_heading(d,'Data and tenancy',2)
add_p(d,'Supabase PostgreSQL stores the platform data. Modules use separate schemas and migrations so that ownership and change history remain clear. Tenant-aware data is scoped by company identifier and protected through application-level authorisation. Private documents, profile photos and candidate documents are stored in Supabase Storage; access is supplied through controlled endpoints and signed URLs rather than public buckets.')
add_heading(d,'Authentication and authorisation',2)
add_p(d,'Supabase Auth provides login and token services. The application is responsible for tenant resolution, roles, permissions and the authorisation checks that protect each operation. Production configuration is supplied through environment variables and secret stores; credentials, connection strings and service tokens are not committed to the repository.')
add_heading(d,'Health monitoring and recovery',2)
add_p(d,'Every production service exposes /alive for basic liveness and /health/ready for dependency-aware readiness. Readiness treats database and authentication dependencies as critical and exposes only a minimal anonymous response; diagnostic detail is protected. Startup migrations run before normal traffic and recurring-job registration. A failed required migration leaves the API non-ready and prevents normal request handling and job registration.')
add_p(d,'The availability target is supported by Railway deployment verification, provider-managed backup facilities and documented recovery procedures. Database and storage recovery, encryption-key management, restore drills, monitoring and deployment recovery are covered by the repository runbooks.')
add_heading(d,'Backup policy',2)
add_p(d,'The operational policy relies on Supabase-managed database and storage backup capabilities, with recovery objectives and restore drills documented in the backup and disaster-recovery runbook. Railway deployment history provides a code rollback mechanism, but application rollback never restores data. Schema changes must use the expand-migrate-contract approach so an older application can coexist briefly with a newer schema during a rollout.')

add_heading(d,'Version Control')
add_p(d,'The One Big Team repository is the source of truth for application code, infrastructure conventions, tests, deployment workflow definitions and runbooks. Changes are made through pull requests. The main branch is protected by the required ci-success status check, with branches required to be current and review recommended before merge.')
add_p(d,'Continuous integration creates an immutable version from the GitHub run number and commit SHA, executes automated quality gates and publishes a deployment artefact. Railway separately builds a service image from the exact deployed commit. Release identity checks use the Railway deployment identifier as the immutable serving identity; release SHA and version values are display labels only.')

add_heading(d,'Appendix A Cloudflare Setup')
add_p(d,'Cloudflare is the DNS and edge entry point for public One Big Team domains. Configure a separate Cloudflare zone for the production domain and restrict access to the account to the people responsible for domain and incident operations.')
table(d,['Step','Action','Acceptance check'],[['1','Add the One Big Team domain to Cloudflare and change the registrar nameservers to the values Cloudflare provides.','The zone is active in Cloudflare.'],['2','In Railway, add each required custom domain to the relevant service. Copy the CNAME and TXT values Railway provides.','Railway shows the domain as awaiting DNS configuration.'],['3','In Cloudflare DNS, create the CNAME pointing to Railway and the TXT ownership-verification record exactly as shown.','Both records resolve; Railway verifies the domain.'],['4','For normal first-level web hosts, enable the Cloudflare proxy (orange cloud). For Railway verification records, leave TXT records DNS-only.','Railway reports the custom domain as verified.'],['5','Set SSL/TLS mode to Full and enable Universal SSL. Do not select Full Strict for a proxied Railway custom domain unless the Railway configuration explicitly supports it.','HTTPS loads without redirect loops.'],['6','Create a permanent redirect for www to the chosen canonical host, preserving path and query string.','Both addresses reach the canonical HTTPS URL.'],['7','Add WAF/rate-limit rules cautiously and test sign-in, API calls and webhooks after every rule change.','Core journeys and health endpoints behave as intended.']],[.45,4.25,2.4])
add_p(d,'Operational note: do not proxy CNAME records used only to prove ownership. A hostname deeper than a first-level subdomain may need Cloudflare proxying disabled unless the account includes Advanced Certificate Manager. Record the chosen domain-to-service mapping in the environment runbook.')

add_heading(d,'Appendix B Postmark Setup')
add_p(d,'Postmark delivers transactional mail. It is used for invitations, password-reset messages, notifications, reminders and controlled contact-form messages. The application sends through the infrastructure email abstraction, so a module never holds or uses a Postmark token directly.')
table(d,['Step','Action','Acceptance check'],[['1','Create a Postmark account and a transactional message stream for the relevant environment. Keep production and non-production streams separate.','A server or stream exists for each intended environment.'],['2','Add and verify the organisation sending domain or sender signature. Use a real monitored sender address.','Postmark marks the sender/domain as confirmed.'],['3','Copy the SPF and DKIM DNS values from Postmark into Cloudflare. Keep the DNS record types and names exactly as supplied.','Postmark reports SPF/DKIM as verified.'],['4','Create the API token with the least scope required for the stream. Store it only as the environment-specific Postmark configuration secret in Railway and GitHub where needed.','No token appears in source control or logs.'],['5','Set the configured From address, reply-to policy and recipient safeguards. Configure approved templates if the deployment uses Postmark templates.','A test email reaches a controlled mailbox with expected branding.'],['6','Use Postmark activity and bounce reporting to investigate delivery failures. Rotate a suspected token and update the Railway secret before revoking the old one.','Delivery, bounce and suppression handling are documented.']],[.45,4.25,2.4])
add_p(d,'DNS records for Postmark, including SPF, DKIM and any verification CNAMEs, are DNS records rather than web traffic. They must remain DNS-only in Cloudflare. Avoid placing personal data or message bodies in logs while diagnosing delivery issues.')

add_heading(d,'Appendix C Railway Setup')
add_p(d,'Railway hosts the production services and maintains a deployment history for recovery. Each environment has separate configuration, secrets and database connections. The repository deployment workflow uses a project-scoped Railway token for the appropriate environment.')
table(d,['Step','Action','Acceptance check'],[['1','Create a Railway project and environments named Test, Staging and Production. Limit environment access to appropriate operators.','All three environments are visible and access is reviewed.'],['2','Create or connect the required services from the GitHub repository: api, app, marketing and admin. Configure each service root and start behaviour from the repository.','Each service can build from a selected commit.'],['3','Add environment variables and secret references: Supabase URLs/keys, connection strings, Postmark token, encryption keys, allowed admin emails and service-specific configuration.','Services start without secrets in repository files.'],['4','Generate Railway domains for initial testing, then attach verified custom domains through Railway networking. Select the correct target port for each service.','Public endpoints resolve and return HTTPS responses.'],['5','Set Railway service health checking to the application readiness endpoint where supported. Confirm /alive and /health/ready from the public service domain.','Healthy deployment is recognised only after readiness succeeds.'],['6','Create a project token per target environment and store it in the matching GitHub secret. Do not use a broad account token when a project token is sufficient.','A deployment workflow can authenticate only to its target scope.'],['7','Keep serverless/sleep settings compatible with API and Hangfire processing. The API service must remain available to process recurring and queued jobs.','Hangfire health check reports at least one active server.']],[.45,4.25,2.4])
add_p(d,'Deployments are verified, not merely started. The workflow captures the currently serving, rollback-capable Railway deployment for each service, deploys the selected commit, waits for the new deployment identity and readiness, then verifies startup migrations. On a failed verification it attempts a deployment-id rollback only when the release safety declaration allows it; otherwise operators follow the documented manual recovery path.')

add_heading(d,'Appendix D GitHub Actions Jobs')
add_p(d,'GitHub Actions definitions live in .github/workflows. A workflow is an automated process triggered by a pull request, push, manual request or schedule. A job is a sequence of steps running on a GitHub-hosted runner. The following table describes the jobs and operational intent in this repository.')
table(d,['Workflow','Trigger','What it does'],[['ci.yml','Every pull request and push to main','Runs dependency audit, versions the build, builds/tests, integration tests, E2E compilation, package creation and the required ci-success gate.'],['deploy-test.yml','Push to main and manual run','Calls reusable deploy.yml for Test; verifies CI, deploys Railway services and checks release identity, readiness and migrations.'],['deploy-staging.yml','Manual run with a selected ref','Calls reusable deploy.yml for Staging.'],['deploy-production.yml','Manual run and GitHub Environment approval','Calls reusable deploy.yml for Production; required reviewers protect the release.'],['deploy.yml','Reusable workflow','Sets release variables, triggers Railway deploys, verifies the running release and coordinates guarded rollback.'],['deployment-health-check.yml','Reusable and manual','Runs deployment health and startup-migration probes.'],['e2e-nightly.yml','03:00 UTC daily and manual','Runs the full Aspire and Playwright E2E suite outside the PR gate.'],['a11y-nightly.yml','04:00 UTC daily and manual','Runs axe-core, keyboard and accessibility journey coverage.'],['perf-nightly.yml','04:00 UTC daily and manual','Runs representative performance tests and preserves result artefacts.'],['dependabot.yml','Scheduled','Checks configured dependency ecosystems for updates.']],[1.45,1.65,4.0])
add_heading(d,'How a deployment works',2)
bullets(d,['A change is proposed in a pull request. CI must complete successfully before it can be merged to main.','A merge to main automatically deploys to Test. Staging and Production are deliberately manual so the operator selects the intended ref and environment.','The reusable deployment workflow resolves the exact Railway environment and service set before making changes. It injects the release SHA/version display values, deploys each service and checks that the deployment currently serving is the expected Railway deployment.','The deployment must pass /health/ready and the API startup-migration health check. A service that is merely running, or an older healthy deployment, cannot satisfy the identity check.','If deployment verification fails, the workflow follows its guarded recovery process. Destructive or non-backward-compatible schema changes must set release-safety.json to refuse automatic application rollback.'])
add_heading(d,'Operator configuration',2)
bullets(d,['Create GitHub Environments named test, staging and production. Configure at least one required reviewer for production and restrict allowed deployment branches/tags.','Add only environment-scoped secrets and variables. Typical entries are RAILWAY_TOKEN, API_HEALTH_BEARER_TOKEN where applicable, and API_BASE_URL.','Keep workflow dispatch files on the default branch. A user needs repository write access to manually start a workflow.','Review scheduled workflow results in the Actions tab. Scheduled runs use the default branch and can be delayed at busy times; they are quality signals, not the primary production scheduler.'])

add_heading(d,'Appendix E Background Jobs')
add_p(d,'One Big Team uses Hangfire, backed by PostgreSQL, for asynchronous and scheduled work. Hangfire runs inside the API process after required startup migrations have succeeded. This arrangement keeps operational work close to the domain modules while retaining durable queues, retries and a central operational view.')
table(d,['Job type','How it is used','Control'],[['Enqueued jobs','A request commits durable business state, then queues a task such as document virus scanning, email delivery or workflow side effects.','The handler and job must be idempotent because Hangfire can retry after failure.'],['Recurring jobs','Modules register recurring work such as reminder processing, reconciliation, retention and expiry checks.','Cron schedules are registered at API startup only after migrations succeed.'],['Recovery and reconciliation jobs','Detect and repair work left incomplete by transient failures or interrupted processing.','Operate across tenants safely and avoid duplicating completed effects.'],['Maintenance jobs','Remove expired idempotency records, purge eligible artefacts and maintain operational state.','Retention policies and audit requirements apply.']],[1.4,3.4,2.3])
add_p(d,'Failures are not ignored. The background-job audit filter records a structured audit event for an unhandled job exception, including a tenant identifier when the job declares a companyId parameter. Hangfire retries transient failures according to job policy; once failed work remains, the Hangfire health check reports degraded status and administrators can review or retry appropriate jobs through the administrative tooling.')
add_p(d,'The key design rule is at-least-once delivery: a job can run more than once. Jobs therefore make their business effect safe to repeat, use durable status or idempotency records where appropriate, and do not treat a successful enqueue as proof that the external effect has already occurred. Health checks report no-server conditions as unhealthy and failed work as degraded.')

add_heading(d,'References',1)
add_p(d,'Internal source material: specifications/architecture/08-deployment-architecture.md; specifications/runbooks/deployment-pipeline.md; specifications/runbooks/backup-and-disaster-recovery.md; .github/workflows/*.yml; and the application background-job registration and infrastructure code.')
add_p(d,'External operational references: Cloudflare DNS proxy status documentation; Railway Working with Domains and Railway CLI token documentation; Postmark sender-signature documentation; GitHub Actions workflow and trigger documentation. These should be rechecked when a provider changes its console or operational policy.')

# Footer
for section in d.sections:
    fp=section.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rr=fp.add_run('One Big Team Technical Architecture | Internal'); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor(89,89,89)
d.core_properties.title='One Big Team Technical Architecture'; d.core_properties.subject='Architecture, operations and deployment guide'; d.core_properties.author='One Big Team'
d.save(OUT)
print(OUT)
